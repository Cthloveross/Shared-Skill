"""SpreadsheetBench corpus types, isolated execution, and workbook scoring.

The answer workbooks are consumed only by :func:`score_workbook`.  Model prompt
construction lives in ``corpus.py`` and receives neither answer paths nor answer
bytes.
"""

from __future__ import annotations

import asyncio
import contextlib
import hashlib
import importlib.util
import io
import json
import os
import re
import shutil
import signal
import subprocess
import tempfile
import time
import uuid
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from typing import Any, Callable, Protocol

REPO_ROOT = Path(__file__).resolve().parents[3]
DEFAULT_EVALUATOR = REPO_ROOT / "reproduction/reference/evaluation/evaluation.py"
DEFAULT_RUNTIME_DOCKERFILE = Path(__file__).with_name("Dockerfile.spreadsheet-runtime")
EXPECTED_RUNTIME_PACKAGES = {
    "numpy": "2.2.6",
    "openpyxl": "3.1.5",
    "pandas": "2.2.3",
    "libreoffice-calc": "4:25.2.3-2+deb13u8",
}
LIBREOFFICE_PROFILE_URI = "file:///tmp/libreoffice-profile"
_RECALC_STATUS = "runtime-recalc.status"
_RECALC_EXIT = "runtime-recalc.exitcode"
_RECALC_STDOUT = "runtime-recalc.stdout.log"
_RECALC_STDERR = "runtime-recalc.stderr.log"
_RECALC_FILES = (_RECALC_STATUS, _RECALC_EXIT, _RECALC_STDOUT, _RECALC_STDERR)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


@dataclass(frozen=True)
class SpreadsheetCase:
    task_id: str
    case_index: int
    input_path: Path
    answer_path: Path

    @property
    def case_id(self) -> str:
        return f"{self.task_id}:case-{self.case_index}"


@dataclass(frozen=True)
class SpreadsheetTask:
    task_id: str
    instruction: str
    instruction_type: str
    answer_position: str
    answer_sheet: str
    data_position: str
    spreadsheet_path: str
    cases: tuple[SpreadsheetCase, ...]


@dataclass(frozen=True)
class ScoreResult:
    passed: bool | None
    score: float | None
    status: str
    message: str
    normalized_answer_position: str


@dataclass(frozen=True)
class ExecutionResult:
    """A single case attempt.

    ``kind`` is one of ``completed``, ``task_failure``, or
    ``infrastructure_error``.  A completed case has a determinate score;
    infrastructure errors remain unknown and are eligible for one retry.
    """

    kind: str
    output_path: Path | None
    returncode: int | None
    stdout: str = ""
    stderr: str = ""
    error: str | None = None
    timed_out: bool = False
    duration_seconds: float = 0.0
    score: ScoreResult | None = None
    recalculation: dict[str, Any] | None = None


def _bounded_text(path: Path, limit: int) -> str:
    if not path.is_file():
        return ""
    return path.read_text(encoding="utf-8", errors="replace")[-limit:]


def _recalculation_provenance(work_dir: Path, limit: int) -> dict[str, Any] | None:
    status_path = work_dir / _RECALC_STATUS
    if not status_path.is_file():
        return None
    raw_exit = _bounded_text(work_dir / _RECALC_EXIT, 100).strip()
    try:
        returncode = int(raw_exit)
    except ValueError:
        returncode = None
    return {
        "attempted": True,
        "status": _bounded_text(status_path, 200).strip() or "unknown",
        "returncode": returncode,
        "stdout": _bounded_text(work_dir / _RECALC_STDOUT, limit),
        "stderr": _bounded_text(work_dir / _RECALC_STDERR, limit),
        "user_installation": LIBREOFFICE_PROFILE_URI,
    }


class CaseExecutor(Protocol):
    async def execute(
        self,
        *,
        task: SpreadsheetTask,
        case: SpreadsheetCase,
        solution: str,
        work_dir: Path,
        timeout_seconds: float,
        attempt: int,
    ) -> ExecutionResult: ...


def _inside(root: Path, candidate: Path) -> bool:
    try:
        candidate.resolve().relative_to(root.resolve())
        return True
    except ValueError:
        return False


def load_tasks(dataset_path: str | Path, *, expected_tasks: int | None = 100) -> list[SpreadsheetTask]:
    """Load the pinned SpreadsheetBench subset and require three cases/task."""

    import json

    dataset_file = Path(dataset_path).resolve()
    dataset_root = dataset_file.parent
    raw = json.loads(dataset_file.read_text(encoding="utf-8"))
    if not isinstance(raw, list):
        raise ValueError("SpreadsheetBench dataset must be a JSON array")
    if expected_tasks is not None and len(raw) != expected_tasks:
        raise ValueError(f"expected {expected_tasks} tasks, found {len(raw)}")

    tasks: list[SpreadsheetTask] = []
    seen: set[str] = set()
    required = {
        "id",
        "instruction",
        "instruction_type",
        "answer_position",
        "answer_sheet",
        "data_position",
        "spreadsheet_path",
    }
    for item in raw:
        if not isinstance(item, dict) or not required.issubset(item):
            missing = sorted(required - set(item if isinstance(item, dict) else {}))
            raise ValueError(f"invalid SpreadsheetBench row; missing {missing}")
        task_id = str(item["id"])
        if task_id in seen:
            raise ValueError(f"duplicate task id: {task_id}")
        seen.add(task_id)
        task_dir = (dataset_root / str(item["spreadsheet_path"])).resolve()
        if not _inside(dataset_root, task_dir):
            raise ValueError(f"task path escapes dataset root: {item['spreadsheet_path']}")
        cases: list[SpreadsheetCase] = []
        for index in (1, 2, 3):
            input_path = task_dir / f"{index}_{task_id}_input.xlsx"
            answer_path = task_dir / f"{index}_{task_id}_answer.xlsx"
            if not input_path.is_file() or not answer_path.is_file():
                raise FileNotFoundError(f"missing case {task_id}:{index} input or answer workbook")
            cases.append(SpreadsheetCase(task_id, index, input_path, answer_path))
        tasks.append(
            SpreadsheetTask(
                task_id=task_id,
                instruction=str(item["instruction"]),
                instruction_type=str(item["instruction_type"]),
                answer_position=str(item["answer_position"]),
                answer_sheet=str(item["answer_sheet"]),
                data_position=str(item["data_position"]),
                spreadsheet_path=str(item["spreadsheet_path"]),
                cases=tuple(cases),
            )
        )
    return tasks


@dataclass(frozen=True)
class _AnswerRange:
    sheet_name: str | None
    cell_range: str


_CELL = re.compile(r"\$?([A-Za-z]{1,3})\$?([1-9][0-9]*)")
_CELL_RANGE = re.compile(
    r"\$?([A-Za-z]{1,3})\$?([1-9][0-9]*):\$?([A-Za-z]{1,3})\$?([1-9][0-9]*)"
)
_SAME_COLUMN_SHORTHAND = re.compile(r"\$?([A-Za-z]{1,3})\$?([1-9][0-9]*):\$?([1-9][0-9]*)")
_WHOLE_COLUMN_RANGE = re.compile(r"\$?([A-Za-z]{1,3}):\$?([A-Za-z]{1,3})")
_MAX_EXCEL_COLUMN = 16_384
_MAX_EXCEL_ROW = 1_048_576


def _column_number(column: str) -> int:
    number = 0
    for character in column.upper():
        number = number * 26 + ord(character) - ord("A") + 1
    return number


def _validate_bounds(start_column: str, start_row: int, end_column: str, end_row: int) -> None:
    start_column_number = _column_number(start_column)
    end_column_number = _column_number(end_column)
    if not (1 <= start_column_number <= end_column_number <= _MAX_EXCEL_COLUMN):
        raise ValueError(f"invalid or reversed Excel columns: {start_column}:{end_column}")
    if not (1 <= start_row <= end_row <= _MAX_EXCEL_ROW):
        raise ValueError(f"invalid or reversed Excel rows: {start_row}:{end_row}")


def _parse_answer_ranges(answer_position: str) -> list[_AnswerRange]:
    parts = [part.strip() for part in answer_position.split(",")]
    if not parts or any(not part for part in parts):
        raise ValueError("answer_position is empty or contains an empty range")

    ranges: list[_AnswerRange] = []
    for part in parts:
        if part.count("!") > 1:
            raise ValueError(f"answer range has multiple sheet separators: {part!r}")
        if "!" not in part:
            if "'" in part:
                raise ValueError(f"unqualified answer range contains a quote: {part!r}")
            ranges.append(_AnswerRange(None, part))
            continue

        raw_sheet, raw_range = part.split("!", 1)
        # SpreadsheetBench v0.1 contains both ``Sheet'!A1`` and
        # ``'Sheet!'A1``.  In both forms the apostrophe adjacent to ``!`` is a
        # misplaced quote, not part of the sheet name or cell range.
        sheet_name = raw_sheet.strip().strip("'").replace("''", "'")
        cell_range = raw_range.strip().strip("'")
        if not sheet_name or not cell_range or "'" in cell_range or "!" in sheet_name:
            raise ValueError(f"unsupported qualified answer range: {part!r}")
        ranges.append(_AnswerRange(sheet_name, cell_range))
    return ranges


def _normalize_cell_range(cell_range: str, *, whole_column_max_row: int | None = None) -> str:
    if match := _CELL.fullmatch(cell_range):
        column, row = match.groups()
        _validate_bounds(column, int(row), column, int(row))
        return f"{column.upper()}{int(row)}"

    if match := _CELL_RANGE.fullmatch(cell_range):
        start_column, start_row, end_column, end_row = match.groups()
        _validate_bounds(start_column, int(start_row), end_column, int(end_row))
        return f"{start_column.upper()}{int(start_row)}:{end_column.upper()}{int(end_row)}"

    if match := _SAME_COLUMN_SHORTHAND.fullmatch(cell_range):
        column, start_row, end_row = match.groups()
        _validate_bounds(column, int(start_row), column, int(end_row))
        return f"{column.upper()}{int(start_row)}:{column.upper()}{int(end_row)}"

    if match := _WHOLE_COLUMN_RANGE.fullmatch(cell_range):
        start_column, end_column = match.groups()
        if whole_column_max_row is None:
            raise ValueError(f"whole-column answer range requires workbook context: {cell_range!r}")
        _validate_bounds(start_column, 1, end_column, whole_column_max_row)
        return f"{start_column.upper()}1:{end_column.upper()}{whole_column_max_row}"

    raise ValueError(f"unsupported answer range syntax: {cell_range!r}")


def _format_answer_range(answer_range: _AnswerRange, normalized_cell_range: str) -> str:
    if answer_range.sheet_name is None:
        return normalized_cell_range
    escaped_sheet = answer_range.sheet_name.replace("'", "''")
    return f"'{escaped_sheet}'!{normalized_cell_range}"


def normalize_answer_position(answer_position: str, answer_sheet: str) -> str:
    """Normalize ranges that do not require workbook dimensions.

    ``answer_sheet`` is retained as source metadata but does not select a
    worksheet.  The official SpreadsheetBench evaluator selects the answer
    workbook's first sheet for unqualified ranges, including records whose
    ``answer_sheet`` field disagrees with the workbook.
    """

    del answer_sheet
    return ",".join(
        _format_answer_range(answer_range, _normalize_cell_range(answer_range.cell_range))
        for answer_range in _parse_answer_ranges(answer_position)
    )


def resolve_answer_position(
    answer_position: str,
    answer_sheet: str,
    *,
    answer_path: Path,
    output_path: Path,
) -> str:
    """Resolve an answer position against the exact scored workbooks.

    Whole-column ranges are bounded to the greatest used row on the relevant
    sheet across the answer and generated workbooks.  Other ranges never need
    workbook access.
    """

    del answer_sheet
    answer_ranges = _parse_answer_ranges(answer_position)
    if not any(_WHOLE_COLUMN_RANGE.fullmatch(item.cell_range) for item in answer_ranges):
        return ",".join(
            _format_answer_range(item, _normalize_cell_range(item.cell_range)) for item in answer_ranges
        )

    import openpyxl

    answer_workbook = openpyxl.load_workbook(answer_path, data_only=True, read_only=True)
    output_workbook = None
    try:
        if output_path.is_file():
            try:
                if output_path.resolve() == answer_path.resolve():
                    output_workbook = answer_workbook
                else:
                    output_workbook = openpyxl.load_workbook(output_path, data_only=True, read_only=True)
            except Exception:
                # Let the official comparator classify an unreadable generated
                # workbook as a task failure.  The answer workbook still gives
                # a non-vacuous bound for the range.
                output_workbook = None

        if not answer_workbook.sheetnames:
            raise ValueError("answer workbook has no worksheets")
        first_sheet = answer_workbook.sheetnames[0]
        normalized: list[str] = []
        for item in answer_ranges:
            max_row = None
            if _WHOLE_COLUMN_RANGE.fullmatch(item.cell_range):
                sheet_name = item.sheet_name or first_sheet
                if sheet_name not in answer_workbook:
                    raise ValueError(f"answer worksheet not found for whole-column range: {sheet_name!r}")
                rows = [answer_workbook[sheet_name].max_row or 1]
                if output_workbook is not None and sheet_name in output_workbook:
                    rows.append(output_workbook[sheet_name].max_row or 1)
                max_row = max(1, *rows)
            normalized.append(
                _format_answer_range(
                    item,
                    _normalize_cell_range(item.cell_range, whole_column_max_row=max_row),
                )
            )
        return ",".join(normalized)
    finally:
        if output_workbook is not None and output_workbook is not answer_workbook:
            output_workbook.close()
        answer_workbook.close()


@lru_cache(maxsize=4)
def _official_comparator(evaluator_path: str) -> Callable[..., tuple[bool, str]]:
    path = Path(evaluator_path)
    spec = importlib.util.spec_from_file_location("spreadsheetbench_official_evaluation", path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"cannot load evaluator: {path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module.compare_workbooks


def score_workbook(
    *,
    task: SpreadsheetTask,
    case: SpreadsheetCase,
    output_path: Path,
    evaluator_path: str | Path = DEFAULT_EVALUATOR,
    comparator: Callable[..., tuple[bool, str]] | None = None,
) -> ScoreResult:
    """Score the generated output, never the input workbook."""

    try:
        normalized = resolve_answer_position(
            task.answer_position,
            task.answer_sheet,
            answer_path=case.answer_path,
            output_path=output_path,
        )
    except Exception as exc:
        return ScoreResult(
            None,
            None,
            "infrastructure_error",
            f"grader range resolution failed: {type(exc).__name__}: {exc}",
            task.answer_position,
        )
    if not output_path.is_file():
        return ScoreResult(False, 0.0, "task_failure", "output workbook is missing", normalized)
    compare = comparator or _official_comparator(str(Path(evaluator_path).resolve()))
    try:
        with contextlib.redirect_stdout(io.StringIO()):
            passed, message = compare(
                str(case.answer_path),
                str(output_path),
                task.instruction_type,
                normalized,
            )
    except Exception as exc:  # grader/runtime failure is unknown, not a zero
        return ScoreResult(
            None, None, "infrastructure_error", f"grader failed: {type(exc).__name__}: {exc}", normalized
        )
    return ScoreResult(bool(passed), float(bool(passed)), "completed", str(message or ""), normalized)


class FakeCaseExecutor:
    """Explicit test backend; produces determinate zero-score trajectories."""

    backend = "fake"

    async def execute(
        self,
        *,
        task: SpreadsheetTask,
        case: SpreadsheetCase,
        solution: str,
        work_dir: Path,
        timeout_seconds: float,
        attempt: int,
    ) -> ExecutionResult:
        del task, solution, timeout_seconds, attempt
        started = time.monotonic()
        work_dir.mkdir(parents=True, exist_ok=True)
        output = work_dir / "output.xlsx"
        shutil.copy2(case.input_path, output)
        score = ScoreResult(False, 0.0, "completed", "explicit fake backend", "fake")
        return ExecutionResult(
            kind="completed",
            output_path=output,
            returncode=0,
            stdout="fake execution",
            duration_seconds=time.monotonic() - started,
            score=score,
        )


class DockerCaseExecutor:
    """Run generated code in a no-network, read-only Docker sandbox."""

    backend = "docker"

    def __init__(
        self,
        *,
        image: str,
        evaluator_path: str | Path = DEFAULT_EVALUATOR,
        memory: str = "2g",
        cpus: float = 1.0,
        stdout_limit: int = 20_000,
        require_recalculation: bool = True,
    ) -> None:
        if not image:
            raise ValueError("runtime.image is required for the Docker backend")
        self.image = image
        self.evaluator_path = Path(evaluator_path)
        self.memory = memory
        self.cpus = cpus
        self.stdout_limit = stdout_limit
        self.require_recalculation = require_recalculation

    async def execute(
        self,
        *,
        task: SpreadsheetTask,
        case: SpreadsheetCase,
        solution: str,
        work_dir: Path,
        timeout_seconds: float,
        attempt: int,
    ) -> ExecutionResult:
        started = time.monotonic()
        work_dir.mkdir(parents=True, exist_ok=True)
        solution_path = work_dir / "solution.py"
        solution_path.write_text(solution, encoding="utf-8")
        output_path = work_dir / "output.xlsx"
        output_path.unlink(missing_ok=True)
        for name in _RECALC_FILES:
            (work_dir / name).unlink(missing_ok=True)

        uid = os.getuid()
        gid = os.getgid()
        recalc = (
            f"""
printf '%s\n' 'started' > /out/{_RECALC_STATUS}
: > /out/{_RECALC_STDOUT}
: > /out/{_RECALC_STDERR}
if ! command -v soffice >/dev/null 2>&1; then
  printf '%s\n' '127' > /out/{_RECALC_EXIT}
  printf '%s\n' 'soffice_missing' > /out/{_RECALC_STATUS}
  exit 44
fi
mkdir -p /tmp/libreoffice-profile /tmp/recalculated
set +e
soffice -env:UserInstallation={LIBREOFFICE_PROFILE_URI} \
  --headless --nologo --nodefault --nofirststartwizard \
  --convert-to xlsx --outdir /tmp/recalculated /out/output.xlsx \
  > /out/{_RECALC_STDOUT} 2> /out/{_RECALC_STDERR}
recalc_code=$?
set -e
printf '%s\n' "$recalc_code" > /out/{_RECALC_EXIT}
if [ "$recalc_code" -ne 0 ]; then
  printf '%s\n' 'soffice_failed' > /out/{_RECALC_STATUS}
  cat /out/{_RECALC_STDOUT}
  cat /out/{_RECALC_STDERR} >&2
  exit 45
fi
if ! test -s /tmp/recalculated/output.xlsx; then
  printf '%s\n' 'converted_output_missing' > /out/{_RECALC_STATUS}
  cat /out/{_RECALC_STDOUT}
  cat /out/{_RECALC_STDERR} >&2
  exit 43
fi
if ! cp /tmp/recalculated/output.xlsx /out/output.xlsx; then
  printf '%s\n' 'copy_failed' > /out/{_RECALC_STATUS}
  exit 46
fi
printf '%s\n' 'success' > /out/{_RECALC_STATUS}
"""
            if self.require_recalculation
            else ""
        )
        script = f"""
set -eu
python /solution.py
test -s /out/output.xlsx || exit 42
{recalc}
"""
        container_name = (
            "pe-sb-"
            + hashlib.sha256(
                f"{os.getpid()}:{task.task_id}:{case.case_index}:{attempt}:{work_dir}:{uuid.uuid4()}".encode()
            ).hexdigest()[:20]
        )
        command = [
            "docker",
            "run",
            "--rm",
            "--name",
            container_name,
            "--network",
            "none",
            "--read-only",
            "--cap-drop",
            "ALL",
            "--security-opt",
            "no-new-privileges",
            "--pids-limit",
            "256",
            "--memory",
            self.memory,
            "--cpus",
            str(self.cpus),
            "--user",
            f"{uid}:{gid}",
            "--tmpfs",
            "/tmp:rw,nosuid,nodev,size=512m",
            "-e",
            "HOME=/tmp",
            "-e",
            "INPUT_PATH=/input.xlsx",
            "-e",
            "OUTPUT_PATH=/out/output.xlsx",
            "-v",
            f"{case.input_path.resolve()}:/input.xlsx:ro",
            "-v",
            f"{solution_path.resolve()}:/solution.py:ro",
            "-v",
            f"{work_dir.resolve()}:/out:rw",
            self.image,
            "sh",
            "-lc",
            script,
        ]
        try:
            process = await asyncio.create_subprocess_exec(
                *command,
                stdout=asyncio.subprocess.PIPE,
                stderr=asyncio.subprocess.PIPE,
                start_new_session=True,
            )
            try:
                stdout_bytes, stderr_bytes = await asyncio.wait_for(process.communicate(), timeout=timeout_seconds)
            except asyncio.TimeoutError:
                with contextlib.suppress(ProcessLookupError):
                    os.killpg(process.pid, signal.SIGKILL)
                await process.wait()
                await self._force_remove_container(container_name)
                return ExecutionResult(
                    kind="infrastructure_error",
                    output_path=None,
                    returncode=None,
                    error=f"case execution timed out after {timeout_seconds}s",
                    timed_out=True,
                    duration_seconds=time.monotonic() - started,
                    recalculation=_recalculation_provenance(work_dir, self.stdout_limit),
                )
            except asyncio.CancelledError:
                with contextlib.suppress(ProcessLookupError):
                    os.killpg(process.pid, signal.SIGKILL)
                with contextlib.suppress(Exception):
                    await asyncio.wait_for(process.wait(), timeout=5)
                await asyncio.shield(self._force_remove_container(container_name))
                raise
        except (FileNotFoundError, OSError) as exc:
            return ExecutionResult(
                kind="infrastructure_error",
                output_path=None,
                returncode=None,
                error=f"docker invocation failed: {type(exc).__name__}: {exc}",
                duration_seconds=time.monotonic() - started,
            )

        stdout = stdout_bytes.decode("utf-8", errors="replace")[-self.stdout_limit :]
        stderr = stderr_bytes.decode("utf-8", errors="replace")[-self.stdout_limit :]
        recalculation = _recalculation_provenance(work_dir, self.stdout_limit)
        if process.returncode != 0:
            infrastructure = (
                process.returncode == 125
                or "docker daemon" in stderr.casefold()
                or recalculation is not None
            )
            kind = "infrastructure_error" if infrastructure else "task_failure"
            if recalculation is not None:
                error = (
                    "recalculation failed: "
                    f"status={recalculation['status']}, soffice_returncode={recalculation['returncode']}"
                )
            else:
                error = f"sandbox exited {process.returncode}"
            return ExecutionResult(
                kind=kind,
                output_path=output_path if output_path.is_file() else None,
                returncode=process.returncode,
                stdout=stdout,
                stderr=stderr,
                error=error,
                duration_seconds=time.monotonic() - started,
                score=(
                    None
                    if infrastructure
                    else ScoreResult(False, 0.0, "task_failure", f"sandbox exited {process.returncode}", "")
                ),
                recalculation=recalculation,
            )

        score = score_workbook(
            task=task,
            case=case,
            output_path=output_path,
            evaluator_path=self.evaluator_path,
        )
        kind = "infrastructure_error" if score.status == "infrastructure_error" else "completed"
        return ExecutionResult(
            kind=kind,
            output_path=output_path,
            returncode=process.returncode,
            stdout=stdout,
            stderr=stderr,
            error=score.message if kind == "infrastructure_error" else None,
            duration_seconds=time.monotonic() - started,
            score=score,
            recalculation=recalculation,
        )

    @staticmethod
    async def _force_remove_container(container_name: str) -> None:
        """Remove a named container after timeout/cancellation; ignore not-found."""

        try:
            process = await asyncio.create_subprocess_exec(
                "docker",
                "rm",
                "-f",
                container_name,
                stdout=asyncio.subprocess.DEVNULL,
                stderr=asyncio.subprocess.DEVNULL,
            )
            await asyncio.wait_for(process.wait(), timeout=15)
        except (OSError, asyncio.TimeoutError):
            return


def docker_runtime_readiness(*, image: str, require_recalculation: bool = True, timeout: float = 30) -> dict[str, Any]:
    """Probe the configured image without executing experiment cases."""

    if not shutil.which("docker"):
        return {
            "ready": False,
            "error": "docker CLI not found",
            "image": image,
            "build_command": docker_runtime_build_command(image=image),
        }
    check = (
        "import json,numpy,openpyxl,pandas; "
        "print(json.dumps({'numpy':numpy.__version__,'openpyxl':openpyxl.__version__,'pandas':pandas.__version__}))"
    )
    command = ["docker", "run", "--rm", "--network", "none", image, "python", "-c", check]
    try:
        result = subprocess.run(command, capture_output=True, text=True, timeout=timeout, check=False)
    except (OSError, subprocess.TimeoutExpired) as exc:
        return {
            "ready": False,
            "error": f"image probe failed: {type(exc).__name__}: {exc}",
            "image": image,
            "build_command": docker_runtime_build_command(image=image),
        }
    try:
        python_packages = json.loads(result.stdout.strip()) if result.returncode == 0 else {}
    except json.JSONDecodeError:
        python_packages = {}
    expected_python = {key: value for key, value in EXPECTED_RUNTIME_PACKAGES.items() if key != "libreoffice-calc"}
    ready = result.returncode == 0 and python_packages == expected_python
    details: dict[str, Any] = {
        "ready": ready,
        "image": image,
        "python_packages": python_packages,
        "expected_python_packages": expected_python,
        "error": None if ready else (result.stderr.strip() or f"probe exited {result.returncode}"),
        "build_command": docker_runtime_build_command(image=image),
    }
    if ready and require_recalculation:
        office = subprocess.run(
            [
                "docker",
                "run",
                "--rm",
                "--network",
                "none",
                image,
                "sh",
                "-lc",
                "command -v soffice >/dev/null && dpkg-query -W -f='${Version}' libreoffice-calc",
            ],
            capture_output=True,
            text=True,
            timeout=timeout,
            check=False,
        )
        office_version = office.stdout.strip() or None
        details["libreoffice_calc"] = office_version
        details["expected_libreoffice_calc"] = EXPECTED_RUNTIME_PACKAGES["libreoffice-calc"]
        if office.returncode != 0 or office_version != EXPECTED_RUNTIME_PACKAGES["libreoffice-calc"]:
            details.update(ready=False, error="configured image has missing or unexpected LibreOffice version")
        if details["ready"]:
            from openpyxl import Workbook

            with tempfile.TemporaryDirectory(prefix="pe-sb-recalc-readiness-") as directory:
                root = Path(directory)
                input_path = root / "input.xlsx"
                output_dir = root / "out"
                output_dir.mkdir()
                workbook = Workbook()
                workbook.active["A1"] = "runtime readiness"
                workbook.save(input_path)
                workbook.close()
                uid, gid = os.getuid(), os.getgid()
                conversion = subprocess.run(
                    [
                        "docker",
                        "run",
                        "--rm",
                        "--network",
                        "none",
                        "--read-only",
                        "--cap-drop",
                        "ALL",
                        "--security-opt",
                        "no-new-privileges",
                        "--pids-limit",
                        "256",
                        "--memory",
                        "2g",
                        "--cpus",
                        "1.0",
                        "--user",
                        f"{uid}:{gid}",
                        "--tmpfs",
                        "/tmp:rw,nosuid,nodev,size=512m",
                        "-e",
                        "HOME=/tmp",
                        "-v",
                        f"{input_path.resolve()}:/input.xlsx:ro",
                        "-v",
                        f"{output_dir.resolve()}:/out:rw",
                        image,
                        "sh",
                        "-lc",
                        (
                            "mkdir -p /tmp/libreoffice-profile /tmp/recalculated && "
                            f"soffice -env:UserInstallation={LIBREOFFICE_PROFILE_URI} "
                            "--headless --nologo --nodefault --nofirststartwizard "
                            "--convert-to xlsx --outdir /tmp/recalculated /input.xlsx && "
                            "test -s /tmp/recalculated/input.xlsx && "
                            "cp /tmp/recalculated/input.xlsx /out/output.xlsx"
                        ),
                    ],
                    capture_output=True,
                    text=True,
                    timeout=timeout,
                    check=False,
                )
                converted = output_dir / "output.xlsx"
                conversion_ready = conversion.returncode == 0 and converted.is_file() and converted.stat().st_size > 0
                details["recalculation_probe"] = {
                    "ready": conversion_ready,
                    "returncode": conversion.returncode,
                    "output_bytes": converted.stat().st_size if converted.is_file() else None,
                    "stdout": conversion.stdout[-20_000:],
                    "stderr": conversion.stderr[-20_000:],
                    "user_installation": LIBREOFFICE_PROFILE_URI,
                    "read_only_root": True,
                    "network": "none",
                }
                if not conversion_ready:
                    details.update(ready=False, error="LibreOffice failed under the exact sandbox security profile")
    return details


def docker_runtime_build_command(
    *,
    image: str = "spreadsheetbench-runtime:20261020",
    dockerfile: str | Path = DEFAULT_RUNTIME_DOCKERFILE,
) -> str:
    """Return the explicit bootstrap command; readiness never builds implicitly."""

    return f"docker build --pull=false -t {image} -f {Path(dockerfile)} {Path(dockerfile).parent}"


def build_docker_runtime(
    *,
    image: str = "spreadsheetbench-runtime:20261020",
    dockerfile: str | Path = DEFAULT_RUNTIME_DOCKERFILE,
    timeout: float = 1800,
) -> dict[str, Any]:
    """Build the pinned runtime when explicitly requested by the CLI/operator."""

    dockerfile_path = Path(dockerfile).resolve()
    command = [
        "docker",
        "build",
        "--pull=false",
        "-t",
        image,
        "-f",
        str(dockerfile_path),
        str(dockerfile_path.parent),
    ]
    try:
        result = subprocess.run(command, capture_output=True, text=True, timeout=timeout, check=False)
    except (OSError, subprocess.TimeoutExpired) as exc:
        return {"built": False, "image": image, "error": f"{type(exc).__name__}: {exc}"}
    return {
        "built": result.returncode == 0,
        "image": image,
        "returncode": result.returncode,
        "stdout": result.stdout[-20_000:],
        "stderr": result.stderr[-20_000:],
    }
